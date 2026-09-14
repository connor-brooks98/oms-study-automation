"""One synthetic Hub lecture: prepare offline, then explicit bounded HTTP acceptance."""

import argparse
import hashlib
import json
import runpy
import time
from io import BytesIO
from pathlib import Path

import httpx
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

FIXTURE = Path(__file__).resolve().parents[1] / "tests/v2/test_gpt_lecture_acceptance.py"
URL = "http://127.0.0.1:62170"


def prepare(destination: Path) -> dict:
    fixture = runpy.run_path(str(FIXTURE))
    destination.mkdir(mode=0o700)
    book = Workbook()
    sheet = book.active
    sheet.append(["Synthetic 1"])
    sheet.append([1, "Synthetic structures", "Fixture teacher"])
    book.save(destination / "tracker.xlsx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(9), Inches(2)).text = fixture[
        "LECTURE_TEXT"
    ]
    slide.shapes.add_picture(BytesIO(fixture["_figure"]()), Inches(1), Inches(2.5), width=Inches(6))
    deck.save(destination / "fixture.pptx")
    (destination / "fixture.txt").write_text(fixture["LECTURE_TEXT"])
    (destination / "cleaning-prompt.md").write_text(
        "Preserve all supplied teaching facts. Return cleaned text as JSON."
    )
    receipt = {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in destination.iterdir()
        if path.is_file()
    }
    (destination / "manifest.json").write_text(json.dumps(receipt, indent=2))
    return receipt


def resume_item(batch: dict, lectures: list, digest: str, *, transcript: bool = False) -> dict:
    assert len(lectures) == 1 and lectures[0]["topic"] == "Synthetic structures"
    assert batch["kind"] == ("transcripts" if transcript else "slides") and len(batch["items"]) == 1
    item = batch["items"][0]
    assert item["state"] == ("needs_review" if transcript else "quarantined")
    assert item["sha256"] == digest
    assert item["error"] == (
        "Transcript cleaning prompt has not been approved in its current form"
        if transcript
        else "iCloud staging root has not been configured"
    )
    assert item["lecture_id"] == lectures[0]["id"]
    return item


def generate(
    artifacts: Path,
    evidence: Path,
    commit: str,
    resume_slide_batch: str | None = None,
    resume_transcript_batch: str | None = None,
) -> dict:
    fixture = runpy.run_path(str(FIXTURE))
    manifest = json.loads((artifacts / "manifest.json").read_text())
    assert set(manifest) == {"tracker.xlsx", "fixture.pptx", "fixture.txt", "cleaning-prompt.md"}
    assert all(
        hashlib.sha256((artifacts / name).read_bytes()).hexdigest() == digest
        for name, digest in manifest.items()
    )
    evidence.mkdir(mode=0o700)
    receipt = {
        "hub_verified": False,
        "native_mac_office_verified": False,
        "synthetic_office_converter": True,
        "production_codex_client": True,
        "automatic_resubmission": False,
        "provider_turn_budget": 2,
    }
    deadline = time.monotonic() + 480

    def save(name, payload):
        (evidence / name).write_text(json.dumps(payload, indent=2))

    with httpx.Client(base_url=URL, timeout=20, follow_redirects=False, trust_env=False) as client:

        def checked(method, path, **kwargs):
            response = client.request(method, path, **kwargs)
            assert response.status_code in {200, 202, 303}, (
                path,
                response.status_code,
                response.text,
            )
            return response

        try:
            health = checked("GET", "/health").json()
            assert health["build_revision"] == commit and health["status"] == "ok"
            checked("GET", "/settings")
            client.headers["X-CSRF-Token"] = client.cookies["study_hub_csrf"]
            lectures = checked("GET", "/api/lectures").json()
            resumed_item = None
            resume_batch = resume_transcript_batch or resume_slide_batch
            if resume_batch:
                prior = checked("GET", f"/api/upload-batches/{resume_batch}").json()
                resumed_item = resume_item(
                    prior,
                    lectures,
                    manifest["fixture.txt" if resume_transcript_batch else "fixture.pptx"],
                    transcript=bool(resume_transcript_batch),
                )
                save("prior-batch.json", prior)
            else:
                assert lectures == []
            status = checked("POST", "/settings/generation/codex/status").json()
            assert status["state"] == "connected" and status["selected_model"] == "gpt-5.5"
            assert "gpt-5.5" in status["image_model_ids"]
            save("preflight.json", {"health": health, "status": status})
            prompt = str(artifacts / "cleaning-prompt.md")
            checked("POST", "/settings/generation/prompts/transcript", json={"path": prompt})
            inspected = checked("POST", "/settings/generation/prompts/transcript/test").json()
            assert (
                inspected["state"] == "valid"
                and inspected["sha256"] == manifest["cleaning-prompt.md"]
            )
            if not resume_batch:
                preview = checked(
                    "POST",
                    "/settings/tracker/preview",
                    headers={"Accept": "application/json"},
                    files={"workbook": ("tracker.xlsx", (artifacts / "tracker.xlsx").read_bytes())},
                ).json()
                assert len(preview["new"]) == 1 and not preview["issues"] and not preview["missing"]
                checked("POST", "/settings/tracker/apply", data={"preview_id": preview["id"]})
            lectures = checked("GET", "/api/lectures").json()
            assert len(lectures) == 1 and lectures[0]["topic"] == "Synthetic structures"
            lecture_id = lectures[0]["id"]
            receipt["lecture_id"] = lecture_id
            save("result.json", receipt)
            for kind, filename in (("slides", "fixture.pptx"), ("transcripts", "fixture.txt")):
                if kind == "slides" and resume_transcript_batch:
                    # Completed in workflow2; queue validates the current source binding.
                    continue
                if resumed_item and kind == (
                    "transcripts" if resume_transcript_batch else "slides"
                ):
                    batch_id = resume_batch
                    checked(
                        "POST",
                        f"/quarantine/{resumed_item['id']}/assign",
                        data={"lecture_id": str(lecture_id)},
                    )
                    receipt["resumed_batch"] = batch_id
                    save("result.json", receipt)
                else:
                    batch_id = checked(
                        "POST",
                        f"/uploads/{kind}",
                        data={"lecture_id": str(lecture_id)},
                        files={"files": (filename, (artifacts / filename).read_bytes())},
                    ).json()["batch_id"]
                confirmed = set()
                while time.monotonic() < deadline:
                    batch = checked("GET", f"/api/upload-batches/{batch_id}").json()
                    save(f"{kind}-batch.json", batch)
                    states = {item["state"] for item in batch["items"]}
                    assert not states & {"failed", "needs_review", "quarantined", "discarded"}, (
                        states
                    )
                    if states == {"complete"}:
                        break
                    for item in batch["items"]:
                        if item["state"] == "awaiting_confirmation" and item["id"] not in confirmed:
                            checked("POST", f"/api/upload-items/{item['id']}/confirm")
                            confirmed.add(item["id"])
                    time.sleep(1)
                else:
                    raise TimeoutError("synthetic ingestion deadline")
            run_id = fixture["_queue"](client, lecture_id)
            receipt["run_id"] = run_id
            save("result.json", receipt)
            while time.monotonic() < deadline:
                state = checked("GET", f"/settings/generation/codex/runs/{run_id}").json()
                save("run-status.json", state)
                assert state["state"] not in {"retrying", "failed", "paused", "interrupted"}, state
                if state["state"] == "awaiting_review":
                    break
                time.sleep(1)
            else:
                raise TimeoutError("synthetic quiz deadline")
            review = checked("GET", f"/studio/runs/{run_id}/review/data").json()
            save("review.json", review)
            assert len(review["questions"]) >= 3
            assert client.post(f"/studio/runs/{run_id}/publication").status_code == 409
            receipt["generated_awaiting_review"] = True
            # Review the actual generated answers before using the existing _verify helper.
        except Exception as error:
            receipt["error"] = f"{type(error).__name__}: {error}"
            raise
        finally:
            save("result.json", receipt)
    return receipt


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--artifacts", type=Path, required=True)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--prepare", action="store_true")
    mode.add_argument("--generate", action="store_true")
    parser.add_argument("--evidence", type=Path)
    parser.add_argument("--expected-commit")
    resume = parser.add_mutually_exclusive_group()
    resume.add_argument("--resume-slide-batch")
    resume.add_argument("--resume-transcript-batch")
    args = parser.parse_args()
    if args.prepare:
        result = prepare(args.artifacts)
    else:
        if args.evidence is None or not args.expected_commit:
            parser.error("--generate requires --evidence and --expected-commit")
        result = generate(
            args.artifacts,
            args.evidence,
            args.expected_commit,
            args.resume_slide_batch,
            args.resume_transcript_batch,
        )
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
