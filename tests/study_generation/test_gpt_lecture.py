import subprocess
import sys
from dataclasses import FrozenInstanceError, replace
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image

from oms_hub.document_processing.assets import persist_asset
from oms_hub.document_processing.domain import (
    DocumentLocator,
    ParsedAsset,
    ParsedDocument,
    ParsedSegment,
    SegmentKind,
    SourceSnapshot,
)
from oms_hub.document_processing.pdf_adapter import PdfProcessor
from oms_hub.document_processing.presentation_render import PresentationRenderer
from oms_hub.document_processing.router import DocumentProcessorRouter, ParserMode
from oms_hub.document_processing.text_adapter import TextProcessor
from oms_hub.study_generation.gpt_lecture import (
    LectureInputs,
    LectureSourceBinding,
    parse_lecture_sources,
    source_asset,
    source_manifest,
    uncovered_objectives,
    validate_lecture_inputs,
)


def test_coverage_is_per_objective_not_question_count():
    assert uncovered_objectives(("LO1", "LO2"), (("LO1",),) * 20) == ("LO2",)
    assert uncovered_objectives(("LO1", "LO2"), (("LO1", "LO2"),)) == ()
    assert uncovered_objectives(("LO3", "LO1", "LO2"), (("LO1", "unknown"),)) == ("LO3", "LO2")


def _inputs(tmp_path: Path, lecture_id: int = 1) -> LectureInputs:
    root = tmp_path / str(lecture_id)
    root.mkdir(parents=True)
    image = BytesIO()
    Image.new("RGB", (20, 20), (lecture_id, 0, 0)).save(image, format="PNG")
    stored = persist_asset(root, "figure-1", "image/png", image.getvalue())
    asset = ParsedAsset(
        stored.key,
        stored.path,
        stored.media_type,
        stored.sha256,
        DocumentLocator("slide 2 image 1", slide_number=2),
        stored.width,
        stored.height,
        "embedded-pptx-image",
    )
    bindings = []
    documents = []
    for role, revision in (("slides", 11), ("cleaned_transcript", 12)):
        path = root / f"{role}.txt"
        path.write_text(f"Synthetic lecture {lecture_id}: {role}")
        snapshot = SourceSnapshot(
            f"{lecture_id}-{role}", role, path, "text/plain", sha256(path.read_bytes()).hexdigest()
        )
        bindings.append(
            LectureSourceBinding(lecture_id, "Heme", 3, revision, role, snapshot, True, True)
        )
        segments = (
            ParsedSegment(
                "block-1",
                SegmentKind.PARAGRAPH,
                f"Evidence for {role}",
                DocumentLocator("slide 2", slide_number=2),
                (asset.key,) if role == "slides" else (),
                style_metadata=("color #FF0000: high yield",),
            ),
        )
        documents.append(
            ParsedDocument(
                snapshot.id,
                snapshot.sha256,
                "txt",
                "synthetic",
                "1",
                segments,
                (asset,) if role == "slides" else (),
                (),
            )
        )
    return LectureInputs(
        lecture_id,
        "Heme",
        3,
        11,
        12,
        bindings[0].snapshot.id,
        bindings[1].snapshot.id,
        (("LO2", "Compare two entities.\nKeep all objective text."), ("LO1", "Explain mechanism.")),
        tuple(documents),
        tuple(bindings),
    )


def test_manifest_retains_verbatim_objectives_styles_and_numbered_assets(tmp_path):
    inputs = _inputs(tmp_path)
    manifest = source_manifest(inputs)
    assert manifest["objectives"][0] == {
        "id": "LO2",
        "text": "Compare two entities.\nKeep all objective text.",
        "sha256": sha256(b"Compare two entities.\nKeep all objective text.").hexdigest(),
    }
    slides, transcript = manifest["sources"]
    assert (slides["source_id"], slides["revision_id"], slides["role"]) == (
        "1-slides",
        11,
        "slides",
    )
    assert transcript["role"] == "cleaned_transcript"
    assert slides["segments"][0]["style_metadata"] == ["color #FF0000: high yield"]
    assert slides["assets"][0]["locator"]["slide_number"] == 2
    assert slides["assets"][0]["sha256"] == inputs.documents[0].assets[0].sha256
    assert source_manifest(inputs) == manifest
    manifest["objectives"][0]["text"] = "mutated output"
    assert source_manifest(inputs)["objectives"][0]["text"].startswith("Compare")
    with pytest.raises(FrozenInstanceError):
        inputs.subject = "Neuro"
    assert source_manifest(replace(inputs, prompt_version="v2"))["sha256"] != manifest["sha256"]
    assert (
        source_manifest(replace(inputs, objectives=(("LO2", "Other objective"),)))["sha256"]
        != manifest["sha256"]
    )


@pytest.mark.parametrize(
    "change",
    [
        {"objectives": ()},
        {"objectives": ((" ", "text"),)},
        {"objectives": (("x", "text"), ("x", "other"))},
        {"objectives": (("x", " \n"),)},
        {"subject": ""},
        {"lecture_id": 0},
        {"exam_number": 0},
        {"bindings": ()},
        {"slide_revision_id": 99},
        {"transcript_revision_id": 99},
        {"slide_source_id": "other-lecture"},
        {"transcript_source_id": "other-transcript"},
        {"documents": ()},
        {"prompt_version": ""},
    ],
)
def test_invalid_selection_is_rejected(tmp_path, change):
    with pytest.raises(ValueError):
        validate_lecture_inputs(replace(_inputs(tmp_path), **change))


@pytest.mark.parametrize(
    "change",
    [
        {"lecture_id": 2},
        {"subject": "Neuro"},
        {"exam_number": 2},
        {"role": "raw_transcript"},
        {"is_current": False},
        {"is_approved": False},
        {"revision_id": 99},
    ],
)
def test_trusted_revision_binding_must_match_scope_and_cleaned_approval(tmp_path, change):
    inputs = _inputs(tmp_path)
    with pytest.raises(ValueError):
        validate_lecture_inputs(
            replace(inputs, bindings=(inputs.bindings[0], replace(inputs.bindings[1], **change)))
        )


@pytest.mark.parametrize(
    "change",
    [
        {"source_sha256": "0" * 64},
        {"source_id": "other-lecture"},
        {"warnings": ("BLOCKER: OCR needed",)},
        {"segments": ()},
        {"parser_name": ""},
        {"parser_version": ""},
    ],
)
def test_parser_evidence_must_be_complete_and_hash_bound(tmp_path, change):
    inputs = _inputs(tmp_path)
    with pytest.raises(ValueError):
        validate_lecture_inputs(
            replace(inputs, documents=(replace(inputs.documents[0], **change), inputs.documents[1]))
        )


def test_source_and_asset_bytes_are_rechecked(tmp_path):
    inputs = _inputs(tmp_path)
    inputs.bindings[1].snapshot.path.write_text("Changed cleaned revision bytes")
    with pytest.raises(ValueError, match="hash"):
        source_manifest(inputs)
    inputs = _inputs(tmp_path, 2)
    asset = inputs.documents[0].assets[0]
    asset.path.write_bytes(b"Changed image bytes")
    with pytest.raises(ValueError, match="hash"):
        source_asset(inputs, inputs.slide_source_id, asset.key)


def test_missing_asset_or_numbered_locator_is_rejected(tmp_path):
    inputs = _inputs(tmp_path)
    asset = inputs.documents[0].assets[0]
    for changed in (
        replace(asset, path=None),
        replace(asset, locator=DocumentLocator("guessed description")),
    ):
        with pytest.raises(ValueError):
            validate_lecture_inputs(
                replace(
                    inputs,
                    documents=(
                        replace(inputs.documents[0], assets=(changed,)),
                        inputs.documents[1],
                    ),
                )
            )
    asset.path.unlink()
    with pytest.raises(ValueError):
        validate_lecture_inputs(inputs)


def test_asset_lookup_qualifies_source_and_rejects_other_lecture_and_transcript(tmp_path):
    first, second = _inputs(tmp_path, 1), _inputs(tmp_path, 2)
    first_asset = source_asset(first, first.slide_source_id, "figure-1")
    second_asset = source_asset(second, second.slide_source_id, "figure-1")
    assert first_asset.sha256 != second_asset.sha256
    for source_id in (second.slide_source_id, first.transcript_source_id):
        with pytest.raises(ValueError):
            source_asset(first, source_id, "figure-1")
    with pytest.raises(ValueError):
        source_asset(first, first.slide_source_id, "invented-image")


def test_image_required_rejects_unavailable_evidence(tmp_path):
    inputs = _inputs(tmp_path)
    slide = replace(
        inputs.documents[0],
        assets=(),
        segments=(replace(inputs.documents[0].segments[0], asset_keys=()),),
    )
    inputs = replace(inputs, documents=(slide, inputs.documents[1]))
    with pytest.raises(ValueError, match="image"):
        validate_lecture_inputs(inputs)
    validate_lecture_inputs(replace(inputs, image_required=False))


class _NoOffice:
    def convert(self, source, destination):
        raise AssertionError("PDF must never invoke Office")


def _vector_pdf_inputs(tmp_path, pages=2, width=500, height=300, logo=False):
    import fitz

    inputs = _inputs(tmp_path)
    path = tmp_path / "vector.pdf"
    with fitz.open() as document:
        for number in range(pages):
            page = document.new_page(width=width, height=height)
            page.insert_text((30, 30), f"Synthetic mechanism on page {number + 1}")
            page.draw_rect((50, 50, 100, 100), color=(1, 0, 0), fill=(1, 0, 0))
            if logo:
                page.insert_text(
                    (30, 150), "Synthetic source text explains the red vector diagram."
                )
                page.insert_text((30, 170), "The tiny blue corner logo is not the teaching figure.")
                raster = BytesIO()
                Image.new("RGB", (10, 10), "blue").save(raster, format="PNG")
                page.insert_image((200, 50, 210, 60), stream=raster.getvalue())
        document.save(path)
    snapshot = replace(
        inputs.bindings[0].snapshot,
        path=path,
        media_type="application/pdf",
        sha256=sha256(path.read_bytes()).hexdigest(),
    )
    return replace(
        inputs,
        documents=(),
        bindings=(replace(inputs.bindings[0], snapshot=snapshot), inputs.bindings[1]),
    )


def _router():
    return DocumentProcessorRouter(
        TextProcessor(), (PdfProcessor(), TextProcessor()), ParserMode.LEGACY
    )


def test_vector_only_pdf_renders_numbered_real_evidence_and_serializable_manifest(
    tmp_path, request
):
    if _isolated_native_check(request):
        return
    import json

    inputs = _vector_pdf_inputs(tmp_path)
    parsed = parse_lecture_sources(
        inputs, _router(), tmp_path / "parsed", renderer=PresentationRenderer(_NoOffice())
    )
    assert inputs.documents == ()
    assert [asset.locator.page_number for asset in parsed.documents[0].assets] == [1, 2]
    assert all(asset.locator.slide_number is None for asset in parsed.documents[0].assets)
    assert all(asset.origin == "full-page-render" for asset in parsed.documents[0].assets)
    assert all(asset.path.is_file() for asset in parsed.documents[0].assets)
    assert "Synthetic mechanism on page 2" in parsed.documents[0].segments[1].text
    manifest = json.loads(json.dumps(source_manifest(parsed)))
    assert Path(manifest["sources"][0]["snapshot"]["path"]).is_file()
    assert Path(manifest["sources"][0]["assets"][0]["path"]).is_file()
    assert manifest["sources"][0]["snapshot"]["sha256"] == inputs.bindings[0].snapshot.sha256


def test_unavailable_vector_render_remains_blocked(tmp_path, request):
    if _isolated_native_check(request):
        return
    inputs = _vector_pdf_inputs(tmp_path)
    with pytest.raises(ValueError, match="image"):
        parse_lecture_sources(inputs, _router(), tmp_path / "parsed")


def test_mixed_vector_pdf_and_tiny_logo_retains_full_page_evidence(tmp_path, request, monkeypatch):
    if _isolated_native_check(request):
        return
    from oms_hub.files.pdf import PdfInspection

    # Isolate visual retention from optional OCR heuristics; text/images/rendering are real.
    monkeypatch.setattr(
        "oms_hub.document_processing.pdf_adapter.inspect_pdf",
        lambda path: PdfInspection("text_based", 1.0, 1, (), False),
    )
    inputs = _vector_pdf_inputs(tmp_path, pages=1, logo=True)
    parsed = parse_lecture_sources(
        inputs, _router(), tmp_path / "parsed", renderer=PresentationRenderer(_NoOffice())
    )
    assets = parsed.documents[0].assets
    assert any((asset.width, asset.height) == (10, 10) for asset in assets)
    full_page = [asset for asset in assets if asset.origin == "full-page-render"]
    assert len(full_page) == 1
    assert (full_page[0].width, full_page[0].height) == (500, 300)
    assert full_page[0].locator.page_number == 1
    with Image.open(full_page[0].path) as raster:
        assert raster.getpixel((75, 75)) == (255, 0, 0)
    manifest = source_manifest(parsed)
    assert len(manifest["sources"][0]["assets"]) == 2


def test_embedded_image_does_not_suppress_render_with_same_slide_locator(tmp_path):
    from oms_hub.document_processing.presentation_render import PresentationRenderResult

    inputs = _inputs(tmp_path)
    embedded = inputs.documents[0].assets[0]
    rendered = replace(embedded, key="slide-2-render", origin="full-slide-render")

    class Router:
        def parse(self, snapshot, root):
            return next(
                document for document in inputs.documents if document.source_id == snapshot.id
            )

    class Renderer:
        def render(self, snapshot, root, **kwargs):
            return PresentationRenderResult((rendered,), ())

    parsed = parse_lecture_sources(inputs, Router(), tmp_path / "parsed", renderer=Renderer())
    assert parsed.documents[0].assets == (embedded, rendered)


@pytest.mark.parametrize(
    "limits,dimensions",
    [
        ({"max_pages": 1, "max_pixels": 1000000}, (500, 300)),
        ({"max_pages": 10, "max_pixels": 10000}, (500, 300)),
    ],
)
def test_renderer_rejects_over_limit_pdf_without_partial_evidence(
    tmp_path, limits, dimensions, request
):
    if _isolated_native_check(request):
        return
    inputs = _vector_pdf_inputs(tmp_path, width=dimensions[0], height=dimensions[1])
    rendered = PresentationRenderer(_NoOffice()).render(
        inputs.bindings[0].snapshot, tmp_path / "renders", **limits
    )
    assert rendered.assets == ()
    assert any("limit" in warning for warning in rendered.warnings)
    assert not list((tmp_path / "renders").glob("*.png"))


def test_parsing_rechecks_snapshot_after_processor_returns(tmp_path):
    inputs = _inputs(tmp_path)

    class MutatingTextProcessor(TextProcessor):
        def parse(self, snapshot, asset_root):
            parsed = super().parse(snapshot, asset_root)
            snapshot.path.write_text("new revision during parse")
            return parsed

    router = DocumentProcessorRouter(TextProcessor(), (MutatingTextProcessor(),), ParserMode.LEGACY)
    with pytest.raises(ValueError, match="hash"):
        parse_lecture_sources(
            replace(inputs, documents=(), image_required=False), router, tmp_path / "assets"
        )


def test_private_manifest_roundtrip_requires_unchanged_source_and_image_bytes(tmp_path):
    from oms_hub.study_generation.gpt_lecture import lecture_inputs_from_manifest

    inputs = _inputs(tmp_path)
    manifest = source_manifest(inputs)
    assert lecture_inputs_from_manifest(manifest) == inputs
    manifest["subject"] = "Neuro"
    with pytest.raises(ValueError, match="digest"):
        lecture_inputs_from_manifest(manifest)
    manifest = source_manifest(inputs)
    inputs.documents[0].assets[0].path.write_bytes(b"changed persisted image")
    with pytest.raises(ValueError, match="hash"):
        lecture_inputs_from_manifest(manifest)


def _isolated_native_check(request):
    # Python 3.13 native import order can crash; exercise the real renderer in a
    # fresh interpreter with PyMuPDF loaded before the application's dependencies.
    if sys.version_info < (3, 13) or "fitz" in sys.modules:
        return False
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            "import fitz, pytest, sys; sys.exit(pytest.main(sys.argv[1:]))",
            request.node.nodeid,
            "-q",
        ],
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    return True


def test_real_pptx_styles_and_embedded_image_survive_private_manifest_reload(tmp_path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor

    from oms_hub.document_processing.anydoc_adapter import AnydocProcessor
    from oms_hub.document_processing.pptx_locator import PptxLocatorEnricher
    from oms_hub.study_generation.gpt_lecture import lecture_inputs_from_manifest
    from tests.document_processing.pptx_factory import SlideFixture, build_pptx

    inputs = _inputs(tmp_path)
    path = build_pptx(
        tmp_path / "lecture.pptx",
        slides=(
            SlideFixture(
                "Lecture title",
                "Explain this synthetic mechanism using the full source evidence here.",
                image=True,
            ),
        ),
    )
    presentation = Presentation(path)
    run = presentation.slides[0].shapes[1].text_frame.paragraphs[0].runs[0]
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.font.bold = True
    presentation.save(path)
    snapshot = replace(
        inputs.bindings[0].snapshot,
        path=path,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        sha256=sha256(path.read_bytes()).hexdigest(),
    )
    inputs = replace(
        inputs,
        documents=(),
        bindings=(replace(inputs.bindings[0], snapshot=snapshot), inputs.bindings[1]),
    )
    router = DocumentProcessorRouter(
        AnydocProcessor(PptxLocatorEnricher()), (TextProcessor(),), ParserMode.ANYDOC
    )
    parsed = parse_lecture_sources(inputs, router, tmp_path / "parsed")
    assert parsed.documents[0].assets[0].locator.slide_number == 1
    assert any(
        "color #FF0000" in style
        for segment in parsed.documents[0].segments
        for style in segment.style_metadata
    )
    sidecar = parsed.run_styles[0]
    assert any(run.resolved_color == "FF0000" and run.bold for run in sidecar.runs)
    assert lecture_inputs_from_manifest(source_manifest(parsed)) == parsed
    with pytest.raises(ValueError, match="run-style"):
        validate_lecture_inputs(replace(parsed, run_styles=()))


def test_image_required_holds_partially_illustrated_slides_when_fallback_unavailable(tmp_path):
    inputs = _inputs(tmp_path)
    slide = inputs.documents[0]
    missing = replace(
        slide.segments[0],
        key="block-2",
        asset_keys=(),
        locator=DocumentLocator("slide 3", slide_number=3),
    )
    inputs = replace(
        inputs, documents=(replace(slide, segments=(*slide.segments, missing)), inputs.documents[1])
    )
    with pytest.raises(ValueError, match="image"):
        validate_lecture_inputs(inputs)
