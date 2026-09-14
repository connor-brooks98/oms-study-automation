import json
from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas

from oms_hub.document_processing.anydoc_adapter import AnydocProcessor
from oms_hub.document_processing.domain import SourceSnapshot
from oms_hub.document_processing.lecture_intake import LECTURE_MEDIA_TYPES
from oms_hub.document_processing.pdf_adapter import PdfProcessor
from oms_hub.document_processing.pptx_locator import PptxLocatorEnricher
from oms_hub.document_processing.router import DocumentProcessorRouter, ParserMode
from oms_hub.document_processing.text_adapter import TextProcessor
from oms_hub.files.atomic import sha256_file
from oms_hub.llm.codex_session import SessionError
from oms_hub.study_generation.gpt_lecture import (
    GeneratedLectureQuiz,
    LectureInputs,
    LectureSourceBinding,
    _batch_sources,
    apply_quiz_instructions,
    generate_lecture_quiz,
    lecture_inputs_from_manifest,
    parse_lecture_sources,
    quiz_instruction_documents,
    source_manifest,
    to_review_drafts,
    validate_generated_quiz,
)
from oms_hub.study_generation.service import _lecture_coverage_targets
from tests.study_generation.test_gpt_lecture import _quiz_payload, _QuizClient


def _source(root: Path, suffix: str) -> Path:
    path = root / ("original" + suffix)
    if suffix == ".pptx":
        presentation = Presentation()
        for number in range(1, 4):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            paragraph = slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(8), Inches(3)
            ).text_frame.paragraphs[0]
            for text, color in (
                (f"Red fact {number}. ", "FF0000"),
                (f"Black excluded {number}.", "000000"),
            ):
                run = paragraph.add_run()
                run.text = text
                run.font.color.rgb = RGBColor.from_string(color)
        presentation.save(path)
    else:
        canvas = Canvas(str(path))
        for number in range(1, 4):
            canvas.drawString(40, 750, f"Page fact {number}.")
            canvas.showPage()
        canvas.save()
    return path


def _router():
    return DocumentProcessorRouter(
        primary=AnydocProcessor(PptxLocatorEnricher()),
        fallbacks=(PdfProcessor(), TextProcessor()),
        mode=ParserMode.ANYDOC,
    )


def _inputs(tmp_path, suffix=".pptx", instructions=""):
    source = _source(tmp_path, suffix)
    transcript = tmp_path / "cleaned.txt"
    transcript.write_text("Spoken excluded context.")
    bindings = tuple(
        LectureSourceBinding(
            1,
            "Neuro",
            1,
            index + 1,
            role,
            SourceSnapshot(role, role, path, LECTURE_MEDIA_TYPES[path.suffix], sha256_file(path)),
            True,
            True,
        )
        for index, (role, path) in enumerate(
            (("slides", source), ("cleaned_transcript", transcript))
        )
    )
    inputs = LectureInputs(
        1,
        "Neuro",
        1,
        1,
        2,
        "slides",
        "cleaned_transcript",
        (("all", "Placeholder"),),
        (),
        bindings=bindings,
        image_required=False,
    )
    parsed = parse_lecture_sources(inputs, _router(), tmp_path / "assets")
    return replace(parsed, instructions=instructions)


@pytest.mark.parametrize(
    "suffix,instructions,unit",
    [
        (".pptx", "only slides 2–3", "slide_number"),
        (".pdf", "Use only pages 2 through 3", "page_number"),
    ],
)
def test_numbered_scope_filters_request_coverage_and_citations_but_keeps_originals(
    tmp_path, suffix, instructions, unit
):
    inputs = apply_quiz_instructions(_inputs(tmp_path, suffix, instructions))
    eligible = quiz_instruction_documents(inputs)
    assert {getattr(segment.locator, unit) for segment in eligible[0].segments} == {2, 3}
    assert eligible[1].segments == ()
    assert len(inputs.documents[0].segments) >= 3
    inputs = replace(inputs, objectives=_lecture_coverage_targets(inputs))
    batches, _ = _batch_sources(inputs)
    raw = batches[0][1]
    assert "excluded context" not in raw
    assert ("Red fact 1" if suffix == ".pptx" else "Page fact 1") not in raw
    assert instructions == json.loads(raw)["quiz_instructions"]
    manifest = source_manifest(inputs)
    assert len(manifest["sources"][0]["segments"]) == len(inputs.documents[0].segments)
    assert source_manifest(lecture_inputs_from_manifest(manifest)) == manifest
    quiz = _payload(inputs)
    validate_generated_quiz(quiz, inputs, require_images=False)
    outside = inputs.documents[0].segments[0].key
    invalid_payload = quiz.model_dump()
    invalid_payload["questions"][0]["source_segments"] = [
        {"source_id": "slides", "segment_key": outside}
    ]
    invalid = GeneratedLectureQuiz.model_validate(invalid_payload)
    with pytest.raises(ValueError):
        validate_generated_quiz(invalid, inputs, require_images=False)


def _payload(inputs):
    payload = _quiz_payload(inputs)
    segment = quiz_instruction_documents(inputs)[0].segments[0]
    for question in payload["questions"]:
        question["source_segments"] = [
            {"source_id": inputs.slide_source_id, "segment_key": segment.key}
        ]
        question["objective_ids"] = [key for key, _text in inputs.objectives]
        question["image"] = None
    return GeneratedLectureQuiz.model_validate(payload)


@pytest.mark.parametrize(
    "instructions",
    [
        "only information in red",
        "only generate a quiz for the info in red",
    ],
)
def test_red_only_uses_actual_runs_never_whole_mixed_color_paragraphs(tmp_path, instructions):
    inputs = apply_quiz_instructions(_inputs(tmp_path, instructions=instructions))
    inputs = replace(inputs, objectives=_lecture_coverage_targets(inputs))
    eligible = quiz_instruction_documents(inputs)
    assert len(eligible[0].segments) == 3
    assert all(segment.key.startswith("red-run-") for segment in eligible[0].segments)
    assert all("Black excluded" not in segment.text for segment in eligible[0].segments)
    assert eligible[0].assets == () and eligible[1].segments == ()
    manifest = source_manifest(inputs)
    assert "Black excluded" in json.dumps(manifest)
    assert source_manifest(lecture_inputs_from_manifest(manifest)) == manifest
    batches, images = _batch_sources(inputs)
    assert not images and "Black excluded" not in batches[0][1]
    assert "Spoken excluded" not in batches[0][1]
    quiz = _payload(inputs)
    drafts = to_review_drafts(quiz, inputs)
    assert drafts[0].source_refs[0].segment_key.startswith("red-run-")
    assert "slide:1:shape:" in drafts[0].source_refs[0].locator


@pytest.mark.parametrize(
    "instructions",
    [
        "Exclude slides 2–3",
        "Do not use slides 2–3",
        "All slides except slides 2–3",
        "only slides 2–3 and only blue text",
        "only slides 2–3 that have red text",
        "only slides 2–3 and only anemia",
        "only slides 0–2",
        "only slides 3–2",
        "only slides 2–8",
        "only slides 1, 3",
        "only slides 1–2 and 3",
        "only highlighted text",
    ],
)
def test_unsupported_inverted_missing_or_mixed_scope_fails_before_queue(tmp_path, instructions):
    with pytest.raises(ValueError):
        apply_quiz_instructions(_inputs(tmp_path, instructions=instructions))


def test_pdf_cannot_infer_red_style_and_unresolved_pptx_colors_fail(tmp_path):
    pdf = apply_quiz_instructions(_inputs(tmp_path, ".pdf"))
    with pytest.raises(ValueError, match="PowerPoint run-color"):
        apply_quiz_instructions(replace(pdf, instructions="only red text"))
    other = tmp_path / "pptx"
    other.mkdir()
    pptx = _inputs(other, instructions="only red text")
    sidecar = pptx.run_styles[0]
    from oms_hub.document_processing.run_styles import StyledTextRunSidecar

    unresolved = StyledTextRunSidecar(
        source_id=sidecar.source_id,
        source_sha256=sidecar.source_sha256,
        parser_version=sidecar.parser_version,
        runs=(sidecar.runs[0].model_copy(update={"resolved_color": None}), *sidecar.runs[1:]),
    )
    with pytest.raises(ValueError, match="unresolved"):
        apply_quiz_instructions(replace(pptx, run_styles=(unresolved,)))


def test_empty_instructions_preserve_legacy_manifest_and_guidance_changes_plan(tmp_path):
    inputs = _inputs(tmp_path)
    assert "instructions" not in source_manifest(inputs)
    assert source_manifest(
        lecture_inputs_from_manifest(source_manifest(inputs))
    ) == source_manifest(inputs)
    ordinary = replace(inputs, instructions="Prefer concise clinical vignettes.")
    assert quiz_instruction_documents(ordinary) == inputs.documents
    assert source_manifest(ordinary)["sha256"] != source_manifest(inputs)["sha256"]
    client = _ScopedClient(tmp_path / "work", ordinary)
    generate_lecture_quiz(
        client,
        "frozen-guidance",
        "gpt-5.5",
        ordinary,
        cancelled=lambda: False,
        on_lifecycle=lambda event: None,
    )
    assert ordinary.instructions in client.requests[0].instructions
    assert json.loads(client.requests[0].source_text)["quiz_instructions"] == ordinary.instructions
    plan = next(client.work_root.rglob("plan.json"))
    original_plan = plan.read_bytes()
    with pytest.raises(SessionError) as error:
        generate_lecture_quiz(
            client,
            "frozen-guidance",
            "gpt-5.5",
            replace(ordinary, instructions="Prefer short explanations."),
            cancelled=lambda: False,
            on_lifecycle=lambda event: None,
        )
    assert error.value.code == "invalid_output"
    assert len(client.requests) == 1 and plan.read_bytes() == original_plan
    with pytest.raises(ValueError, match="4000"):
        replace(inputs, instructions="x" * 4001)


class _ScopedClient(_QuizClient):
    def generate(self, request, **kwargs):
        result = super().generate(request, **kwargs)
        return replace(result, text=_payload(self.inputs).model_dump_json())


def test_red_only_real_parser_queue_fake_generation_review_publication_and_export(tmp_path):
    from types import SimpleNamespace

    from oms_hub.ingestion.domain import UploadKind
    from oms_hub.repositories import CatalogRepository
    from oms_hub.slides.pipeline import SlidePipeline
    from oms_hub.study_generation.gpt_lecture import GptLectureWorker
    from oms_hub.study_generation.practice_review import PracticeReviewService
    from oms_hub.study_generation.quiz_export import export_reviewed_quiz
    from oms_hub.study_generation.quiz_images import StudioQuizImageService
    from oms_hub.study_generation.repository import GenerationRepository
    from oms_hub.study_generation.service import GptLectureService
    from oms_hub.study_generation.studio_repository import StudioRepository
    from oms_hub.transcripts.pipeline import TranscriptPipeline
    from oms_hub.web.gpt_export_routes import _snapshot
    from tests.v2.test_lecture_intake_formats import (
        _Cleaner,
        _Converter,
        _environment,
        _Prompt,
        _stage,
    )

    database, settings, ingestion, staging, _service, lecture = _environment(tmp_path)
    source = _source(tmp_path, ".pptx")
    material = _stage(ingestion, staging, source, UploadKind.SLIDES, lecture)
    SlidePipeline(database, settings, _Converter()).process(material.item_id)
    transcript_path = tmp_path / "spoken.txt"
    transcript_path.write_text("Spoken excluded context.")
    transcript = _stage(ingestion, staging, transcript_path, UploadKind.TRANSCRIPTS, lecture)
    TranscriptPipeline(database, settings, _Prompt(), _Cleaner()).process(transcript.item_id)
    repository = StudioRepository(database)
    service = GptLectureService(
        CatalogRepository(database),
        ingestion,
        repository,
        _router(),
        None,
        settings.data_dir / "sources",
        owner_id="owner",
        model="gpt-5.5",
    )
    with pytest.raises(ValueError, match="automatic coverage"):
        service.queue(
            lecture,
            owner_id="owner",
            label="Incompatible manual scope",
            objectives=(("all", "Cover the whole lecture"),),
            instructions="only red text",
        )
    queued = service.queue(lecture, owner_id="owner", instructions="only red text")
    inputs = service.load_inputs(queued)
    assert not inputs.image_required and len(inputs.objectives) == 3
    client = _ScopedClient(settings.data_dir / "work", inputs)
    images = StudioQuizImageService(repository, settings.data_dir / "media")
    worker = GptLectureWorker(
        repository, client, service.load_inputs, "gpt-5.5", images, settings.data_dir / "evidence"
    )
    run = repository.claim_next_run()
    worker.run(run)
    assert repository.get_run(run.id).state == "awaiting_review"
    assert len(client.requests) == 1
    assert "Black excluded" not in client.requests[0].source_text
    assert "Spoken excluded" not in client.requests[0].source_text
    review = PracticeReviewService(repository, images, lecture_validator=worker.validate_review)
    for question in review.review(run.id):
        assert question.draft.source_refs[0].segment_key.startswith("red-run-")
        review.verify_generated_answer(run.id, question.draft.question_id)
    generation = GenerationRepository(database, practice_review=review)
    published = generation.publish_reviewed_studio_quiz(run.id)
    assert len(published.quiz.questions) == 3
    request = SimpleNamespace(
        app=SimpleNamespace(
            state=SimpleNamespace(
                studio_repository=repository,
                practice_review=review,
                generation_repository=generation,
                settings=settings,
            )
        )
    )
    quiz, paths, provenance, _identity = _snapshot(request, run.id, "owner")
    files = export_reviewed_quiz(quiz, paths, provenance, tmp_path / "export")
    assert files[2].read_bytes().startswith(b"%PDF")
    assert all(
        ref["segment_key"].startswith("red-run-")
        for value in provenance["questions"].values()
        for ref in value["source_refs"]
    )
    assert len(client.requests) == 1
