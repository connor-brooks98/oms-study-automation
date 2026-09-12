"""Provider-free route acceptance through the actual application constructors."""

import hashlib
import json
import os
import socket
import subprocess
import sys
import threading
import time
from dataclasses import asdict
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from oms_hub.app import create_app
from oms_hub.config import Settings
from oms_hub.document_processing.domain import DocumentLocator, ParsedAsset
from oms_hub.document_processing.presentation_render import (
    PresentationRenderer,
    PresentationRenderResult,
)
from oms_hub.llm.codex_session import SessionError, SessionLifecycle, SessionResult, SessionStatus
from oms_hub.repositories import LectureInput
from oms_hub.study_generation.native_quiz import parse_native_quiz
from oms_hub.study_generation.quiz_export import export_reviewed_quiz

LECTURE_TEXT = (
    "Synthetic teaching fixture: structure A is the large blue ring and binds probe A. "
    "Structure B is the small orange disk and binds probe B. Compare the structures "
    "by shape and size, then select the matching laboratory probe."
)
OBJECTIVES = [
    {"id": "lo-1", "text": "Distinguish the two structures in the supplied diagram."},
    {"id": "lo-2", "text": "Select the probe associated with each structure."},
]


def _figure():
    image = Image.new("RGB", (500, 240), "white")
    draw = ImageDraw.Draw(image)
    draw.ellipse((35, 30, 205, 200), outline="#2563eb", width=20)
    draw.ellipse((325, 70, 415, 160), fill="#f97316")
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


class FixtureOfficeConverter:
    """One synthetic slide; does not claim macOS/Windows Office acceptance."""

    def convert(self, source, destination):
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfgen.canvas import Canvas

        assert source.suffix == ".pptx"
        deck = Presentation(source)
        destination.parent.mkdir(parents=True, exist_ok=True)
        canvas = Canvas(str(destination), pagesize=(720, 540))
        for slide in deck.slides:
            text = " ".join(shape.text for shape in slide.shapes if shape.has_text_frame)
            for index, start in enumerate(range(0, len(text), 90)):
                canvas.drawString(25, 500 - index * 16, text[start : start + 90])
            canvas.drawImage(ImageReader(BytesIO(_figure())), 80, 65, width=500, height=240)
            canvas.showPage()
        canvas.save()


class IsolatedFixtureRenderer(PresentationRenderer):
    """Exercise the real rasterizer with native lifetime confined to one process."""

    def _rasterize(self, pdf_path, asset_root, **limits):
        code = """
import fitz, json, sys
from dataclasses import asdict
from pathlib import Path
from oms_hub.document_processing.presentation_render import PresentationRenderer
result = PresentationRenderer(None)._rasterize(Path(sys.argv[1]), Path(sys.argv[2]),
                                              **json.loads(sys.argv[3]))
print(json.dumps(asdict(result), default=str))
"""
        process = subprocess.run(
            [sys.executable, "-c", code, str(pdf_path), str(asset_root), json.dumps(limits)],
            capture_output=True,
            text=True,
            timeout=30,
            check=True,
        )
        payload = json.loads(process.stdout.splitlines()[-1])
        return PresentationRenderResult(
            tuple(
                ParsedAsset(
                    **(
                        asset
                        | {
                            "path": Path(asset["path"]),
                            "locator": DocumentLocator(**asset["locator"]),
                        }
                    )
                )
                for asset in payload["assets"]
            ),
            tuple(payload["warnings"]),
        )


class FixtureSession:
    def __init__(self, work_root):
        self.work_root = work_root
        self.requests = []
        self.invalid_citation = False
        self.blocked = False

    def close(self):
        pass

    def status(self):
        return SessionStatus(
            "unavailable" if self.blocked else "connected",
            ("fixture-model",),
            ("fixture-model",),
            None,
            "capability_unverified" if self.blocked else None,
            account_connected=True,
        )

    def generate(self, request, *, cancelled, on_lifecycle):
        self.requests.append(request)
        if self.blocked:
            raise SessionError("capability_unverified")
        assert not cancelled()
        for phase, thread, turn in (
            ("dispatching", None, None),
            ("thread_created", "fixture-thread", None),
            ("turn_started", "fixture-thread", "fixture-turn"),
        ):
            on_lifecycle(SessionLifecycle(request.request_id, phase, thread, turn))
        if request.request_id.startswith("transcript:"):
            text = json.dumps({"text": request.source_text.split("\n", 1)[1]})
        else:
            assert ":batch-" in request.request_id, "outline/chat work was not requested"
            evidence = json.loads(request.source_text)
            slides = next(source for source in evidence["sources"] if source["role"] == "slides")
            image = evidence["images"][0]
            segment = next(
                segment
                for segment in slides["segments"]
                if segment["locator"]["slide_number"] == image["locator"]["slide_number"]
            )
            for path, digest in zip(request.image_paths, request.image_sha256, strict=True):
                assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
            questions = []
            for index in range(3):
                questions.append(
                    {
                        "id": f"case-{index + 1}",
                        "stem": (
                            f"Synthetic laboratory case {index + 1}: "
                            "which probe binds the large ring?"
                        ),
                        "choices": ["Probe A", "Probe B"],
                        "correct_index": 0,
                        "rationale": "The supplied lecture associates the large ring with probe A.",
                        "distractor_explanations": [
                            "A matches the ring.",
                            "B matches the small disk.",
                        ],
                        "objective_ids": [objective["id"] for objective in evidence["objectives"]],
                        "source_segments": [
                            {
                                "source_id": slides["source_id"],
                                "segment_key": "invented"
                                if self.invalid_citation
                                else segment["key"],
                            }
                        ],
                        "image": {"source_id": slides["source_id"], "asset_key": image["asset_key"]}
                        if index == 0
                        else None,
                    }
                )
            text = json.dumps({"title": "Synthetic lecture quiz", "questions": questions})
        on_lifecycle(
            SessionLifecycle(request.request_id, "completed", "fixture-thread", "fixture-turn")
        )
        return SessionResult("fixture-thread", "fixture-turn", text)


def _build_app(tmp_path, monkeypatch):
    import oms_hub.app as app_module

    session = FixtureSession(tmp_path / "data/codex-work")
    monkeypatch.setattr(app_module, "CodexSessionClient", lambda *args, **kwargs: session)
    monkeypatch.setattr(
        app_module, "SerialOfficeConverter", lambda *args, **kwargs: FixtureOfficeConverter()
    )
    monkeypatch.setattr(app_module, "PresentationRenderer", IsolatedFixtureRenderer)
    prompt = tmp_path / "cleaning-prompt.md"
    prompt.write_text("Preserve all supplied teaching facts. Return cleaned text as JSON.")
    app = create_app(
        Settings(
            _env_file=None,
            data_dir=tmp_path / "data",
            database_url=f"sqlite:///{tmp_path / 'hub.db'}",
            study_root=tmp_path / "study",
            icloud_staging_root=tmp_path / "fixture-staging",
            codex_executable=tmp_path / "fixture-codex",
            codex_binary_sha256="a" * 64,
            codex_model="fixture-model",
            transcript_prompt_path=prompt,
            transcript_prompt_sha256=hashlib.sha256(prompt.read_bytes()).hexdigest(),
        )
    )

    def paid_or_google(*args, **kwargs):
        raise AssertionError("paid API/Google must not be consulted for GPT lecture workflow")

    for name in ("clean", "generate_text", "generate_text_for_task", "for_task", "test_connection"):
        monkeypatch.setattr(app.state.llm_service, name, paid_or_google)
    monkeypatch.setattr(app.state.medical_accuracy_gate, "validate", paid_or_google)
    monkeypatch.setattr(app.state.studio_worker.gateway, "ask_studio", paid_or_google)
    monkeypatch.setattr(app.state.notebook_connection, "invalidate", paid_or_google)
    lecture_id = app.state.catalog_repository.upsert_lecture(
        LectureInput("Heme", 3, 1, "Synthetic structures", "Fixture teacher", None)
    )
    client = TestClient(app)  # Explicit worker ticks; no background lifespan/provider activity.
    assert client.get(f"/lectures/{lecture_id}").status_code == 200
    client.headers["X-CSRF-Token"] = client.cookies["study_hub_csrf"]
    return app, client, session, lecture_id


def _upload_sources(app, client, lecture_id):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(9), Inches(2)).text = LECTURE_TEXT
    slide.shapes.add_picture(BytesIO(_figure()), Inches(1), Inches(2.5), width=Inches(6))
    pptx = BytesIO()
    deck.save(pptx)
    for kind, name, payload in (
        ("slides", "fixture.pptx", pptx.getvalue()),
        ("transcripts", "fixture.txt", LECTURE_TEXT.encode()),
    ):
        response = client.post(
            f"/uploads/{kind}",
            data={"lecture_id": str(lecture_id)},
            files={"files": (name, payload)},
        )
        assert response.status_code == 202, response.text
        batch = client.get(f"/api/upload-batches/{response.json()['batch_id']}").json()
        for item in batch["items"]:
            if item["state"] == "awaiting_confirmation":
                assert client.post(f"/api/upload-items/{item['id']}/confirm").status_code == 200
        assert app.state.ingestion_worker.run_once()
    revisions = app.state.ingestion_repository.list_current_revisions(lecture_id)
    assert len(revisions) == 2 and all(r.state == "current" for r in revisions)
    return revisions


def _queue(client, lecture_id, label="Synthetic image quiz"):
    response = client.post(
        f"/lectures/{lecture_id}/gpt-quiz",
        json={"label": label, "objectives": OBJECTIVES, "require_images": True},
    )
    assert response.status_code == 202, response.text
    return response.json()["run_id"]


def _verify(client, run_id):
    response = client.get(f"/studio/runs/{run_id}/review/data")
    assert response.status_code == 200, response.text
    for question in response.json()["questions"]:
        assert question["verification_required"] and not question["verified_at"]
        checked = client.post(f"/studio/runs/{run_id}/questions/{question['id']}/verify-answer")
        assert checked.status_code == 200, checked.text
    return response.json()


def test_upload_clean_generate_review_publish_and_public_grade_without_google(
    tmp_path, monkeypatch, request
):
    app, client, session, lecture_id = _build_app(tmp_path, monkeypatch)
    try:
        revisions = _upload_sources(app, client, lecture_id)
        assert len(session.requests) == 1 and session.requests[0].request_id.startswith(
            "transcript:"
        )
        assert any(r.provenance_kind == "llm_cleaned" for r in revisions)
        run_id = _queue(client, lecture_id)
        assert app.state.studio_worker.run_once()
        status = client.get(f"/settings/generation/codex/runs/{run_id}")
        assert status.json()["state"] == "awaiting_review", status.text
        assert client.post(f"/studio/runs/{run_id}/publication").status_code == 409
        assert client.get(f"/studio/runs/{run_id}/preview/content").status_code == 409
        _verify(client, run_id)
        preview = client.get(f"/studio/runs/{run_id}/preview/content")
        assert preview.status_code == 200, preview.text
        drafts = app.state.practice_review.review(run_id)
        stored = {
            q.chosen_image.key: app.state.studio_repository.import_review_image(
                run_id, q.chosen_image.key
            )
            for q in drafts
            if q.chosen_image
        }
        publication = client.post(f"/studio/runs/{run_id}/publication")
        assert publication.status_code == 200, publication.text
        token = publication.json()["token"]
        assert client.post(f"/studio/runs/{run_id}/publication").json()["token"] == token
        public = client.get(f"/public/quizzes/{token}/content")
        assert public.status_code == 200
        questions = public.json()["questions"]
        assert len(questions) == 3
        assert all("correct" not in key and key != "rationale" for q in questions for key in q)
        assert all(
            "lo-1" in q["learning_objective"] and "lo-2" in q["learning_objective"]
            for q in questions
        )
        image = next(q["image_url"] for q in questions if q.get("image_url"))
        media = client.get(image)
        assert media.status_code == 200 and media.headers["content-type"].startswith("image/png")
        answer = client.post(
            f"/public/quizzes/{token}/answer",
            json={"question_id": questions[0]["id"], "choice_id": "c1"},
        )
        assert answer.status_code == 200 and answer.json()["correct"] is True
        assert "rationale" in answer.json()
        assert len(session.requests) == 2  # One cleaning turn and one quiz turn; no outline.
        published = app.state.generation_repository.published_quiz(token)
        reviewed = app.state.practice_review.to_native_quiz(run_id, title=published.quiz.title)
        assert reviewed == published.quiz
        drafts = app.state.practice_review.review(run_id)
        provenance = {
            "questions": {
                q.id: {
                    "objective_ids": draft.learning_objective.split(", "),
                    "source_refs": [asdict(ref) for ref in draft.draft.source_refs],
                }
                for q, draft in zip(reviewed.questions, drafts, strict=True)
            },
            "image_sha256": {key: value.sha256 for key, value in stored.items()},
        }
        raw, bundle, pdf = export_reviewed_quiz(
            reviewed,
            {key: value.path for key, value in stored.items()},
            provenance,
            Path(os.environ.get("OMS_B7_EVIDENCE_DIR", str(tmp_path / "exports"))),
        )
        assert parse_native_quiz(raw.read_text()) == published.quiz
        with ZipFile(bundle) as archive:
            manifest = json.loads(archive.read("manifest.json"))
            assert manifest["provenance"] == provenance
            assert archive.read(next(iter(manifest["images"].values()))["path"]) == media.content
        from pypdf import PdfReader

        pdf_pages = [page.extract_text() for page in PdfReader(pdf).pages]
        assert all(len(text.split()) > 8 for text in pdf_pages), "blank export page"
        pdf_text = "\n".join(pdf_pages)
        assert all(q.stem in pdf_text for q in reviewed.questions)
        assert "lo-1" in pdf_text and "lo-2" in pdf_text
    finally:
        client.close()
        app.state.database.close()


@pytest.mark.parametrize("failure", ["source_changed", "missing_image", "invalid_citation"])
def test_lecture_routes_block_invalid_or_missing_evidence(tmp_path, monkeypatch, request, failure):
    app, client, session, lecture_id = _build_app(tmp_path, monkeypatch)
    try:
        revisions = _upload_sources(app, client, lecture_id)
        run_id = _queue(client, lecture_id)
        if failure == "source_changed":
            next(
                r for r in revisions if r.kind.value == "slides"
            ).immutable_source_path.write_bytes(b"Changed synthetic slide bytes")
        session.invalid_citation = failure == "invalid_citation"
        assert app.state.studio_worker.run_once()
        run = app.state.studio_repository.get_run(run_id)
        if failure == "missing_image":
            _verify(client, run_id)
            selected = app.state.practice_review.review(run_id)[0].chosen_image
            app.state.studio_repository.import_review_image(run_id, selected.key).path.unlink()
            assert client.get(f"/studio/runs/{run_id}/preview/content").status_code == 409
        else:
            assert run.state.value == "failed" and run.diagnostic_source == "invalid_output"
        assert client.post(f"/studio/runs/{run_id}/publication").status_code == 409
        assert app.state.studio_repository.get_run(run_id).published_token is None
        assert len(session.requests) == (1 if failure == "source_changed" else 2)
    finally:
        client.close()
        app.state.database.close()


def test_connected_account_does_not_imply_verified_generation_capability(
    tmp_path, monkeypatch, request
):
    app, client, session, lecture_id = _build_app(tmp_path, monkeypatch)
    try:
        _upload_sources(app, client, lecture_id)
        session.blocked = True
        status = client.post("/settings/generation/codex/status")
        assert status.status_code == 200
        assert status.json()["account_connected"] is True
        assert status.json()["error_code"] == "capability_unverified"
        run_id = _queue(client, lecture_id)
        assert app.state.studio_worker.run_once()
        run = client.get(f"/settings/generation/codex/runs/{run_id}").json()
        assert run["state"] == "paused" and run["diagnostic_source"] == "capability_unverified"
        assert (
            app.state.studio_repository.run_artifact(run_id, f"gpt:attempt:{run_id}:batch-0001")
            is None
        )
        assert client.post(f"/studio/runs/{run_id}/publication").status_code == 409
    finally:
        client.close()
        app.state.database.close()


@pytest.mark.skipif(os.environ.get("OMS_B6_BROWSER") != "1", reason="opt-in installed Chrome check")
def test_desktop_mobile_browser_generation_review_publication_and_readiness(tmp_path, monkeypatch):
    import uvicorn
    from playwright.sync_api import expect, sync_playwright

    app, client, session, lecture_id = _build_app(tmp_path, monkeypatch)
    _upload_sources(app, client, lecture_id)
    evidence = Path(os.environ.get("OMS_B6_EVIDENCE_DIR", str(tmp_path / "browser-evidence")))
    evidence.mkdir(parents=True, exist_ok=True)
    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    port = listener.getsockname()[1]
    base = f"http://127.0.0.1:{port}"
    server = uvicorn.Server(uvicorn.Config(app, log_level="error", lifespan="off"))
    thread = threading.Thread(target=lambda: server.run(sockets=[listener]), daemon=True)
    thread.start()
    deadline = time.monotonic() + 10
    while not server.started and time.monotonic() < deadline:
        time.sleep(0.05)
    assert server.started
    observations = []
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(channel="chrome", headless=True)
            for name, width, height in (("desktop", 1365, 900), ("mobile", 390, 844)):
                context = browser.new_context(viewport={"width": width, "height": height})
                # Google fonts are optional; no external service is needed for the workflow.
                context.route(
                    "**/*",
                    lambda route: (
                        route.continue_() if route.request.url.startswith(base) else route.abort()
                    ),
                )
                page = context.new_page()
                errors = []
                page.on("pageerror", lambda error, errors=errors: errors.append(str(error)))

                def capture(stage, page=page, name=name):
                    overflow = page.evaluate(
                        "document.documentElement.scrollWidth > innerWidth + 2"
                    )
                    observations.append({"viewport": name, "stage": stage, "overflow": overflow})
                    page.screenshot(path=str(evidence / f"{name}-{stage}.png"), full_page=True)
                    assert not overflow, f"horizontal overflow at {name} {stage}"

                page.goto(f"{base}/lectures/{lecture_id}")
                page.get_by_label("Quiz title", exact=True).fill(f"Browser {name} lecture quiz")
                page.get_by_label("Learning objectives", exact=True).fill(
                    "\n".join(item["text"] for item in OBJECTIVES)
                )
                page.locator("[data-gpt-lecture]").scroll_into_view_if_needed()
                capture("lecture")
                page.get_by_role("button", name="Create quiz for review", exact=True).click()
                page.get_by_role("link", name="Follow quiz progress", exact=True).click()
                expect(page.locator("[data-gpt-run-state]")).to_have_text("Queued")
                assert app.state.studio_worker.run_once()
                page.get_by_role("button", name="Refresh", exact=True).click()
                expect(page.locator("[data-gpt-run-state]")).to_have_text(
                    "Ready for question review"
                )
                capture("progress")
                page.get_by_role("link", name="Open question review", exact=True).click()
                expect(page.get_by_role("button", name="Verify answer", exact=True)).to_have_count(
                    3
                )
                capture("review")
                for remaining in (2, 1, 0):
                    page.get_by_role("button", name="Verify answer", exact=True).first.click()
                    expect(
                        page.get_by_role("button", name="Verify answer", exact=True)
                    ).to_have_count(remaining)
                page.get_by_role("link", name="Preview and publish", exact=True).click()
                expect(page.locator(".quiz-question-image img")).to_be_visible()
                page.get_by_role("button", name="Publish quiz", exact=True).click()
                page.wait_for_url("**/public/quizzes/*")
                expect(page.locator(".quiz-question-image img")).to_be_visible()
                expect(page.locator(".quiz-question-image img")).to_have_js_property(
                    "complete", True
                )
                assert (
                    page.locator(".quiz-question-image img").evaluate("img => img.naturalWidth") > 0
                )
                capture("public")
                page.locator(".quiz-answer").first.click()
                page.get_by_role("button", name="Submit Answer", exact=True).click()
                expect(page.locator(".quiz-feedback")).to_contain_text("Correct")
                capture("graded")
                session.blocked = True
                page.goto(f"{base}/settings")
                page.get_by_role("button", name="Check connection", exact=True).click()
                expect(page.locator("[data-gpt-status]")).to_have_text(
                    "Account: connected. Generation readiness: capabilities not verified."
                )
                capture("readiness")
                session.blocked = False
                assert not errors, errors
                context.close()
            browser.close()
        (evidence / "observations.json").write_text(json.dumps(observations, indent=2))
    finally:
        server.should_exit = True
        thread.join(timeout=10)
        listener.close()
        client.close()
        app.state.database.close()
    assert not thread.is_alive()
