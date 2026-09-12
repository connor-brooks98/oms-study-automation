"""Native Office smoke coverage for controlled Windows release validation."""

from __future__ import annotations

import hashlib
import json
import sys
from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches
from pypdf import PdfReader

from oms_hub.document_processing.presentation_render import PresentationRenderer
from oms_hub.files.office import SerialOfficeConverter
from oms_hub.files.pdf import validate_pdf
from oms_hub.study_generation.native_quiz import grade_answer, parse_native_quiz
from oms_hub.study_generation.quiz_export import export_reviewed_quiz
from oms_hub.study_generation.quiz_images import sanitize_quiz_image
from tests.document_processing.pptx_factory import SlideFixture, build_pptx
from tests.study_generation.test_quiz_export import export_fixture


@pytest.mark.windows_office
@pytest.mark.skipif(sys.platform != "win32", reason="requires Windows desktop Microsoft Office")
def test_windows_powerpoint_images_and_reviewed_quiz_exports(tmp_path: Path) -> None:
    source = build_pptx(
        tmp_path / "office-smoke.pptx",
        slides=(
            SlideFixture(
                "Embedded image", "A red rectangle appears below this synthetic text.", image=True
            ),
            SlideFixture("Text slide", "This second synthetic slide has distinct visible text."),
        ),
    )
    presentation = Presentation(source)
    picture = next(
        shape
        for shape in presentation.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    picture.width, picture.height = Inches(2), Inches(1.5)
    oval = presentation.slides[1].shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1), Inches(3.5), Inches(2), Inches(1.5)
    )
    oval.fill.solid()
    oval.fill.fore_color.rgb = RGBColor(20, 20, 200)
    presentation.save(source)
    destination = tmp_path / "office-smoke.pdf"

    converter = SerialOfficeConverter(timeout_seconds=120, admission_timeout_seconds=10)
    converter.convert(source, destination)

    assert destination.is_file()
    assert validate_pdf(destination).page_count == 2
    pages = PdfReader(destination).pages
    assert "Embedded image" in pages[0].extract_text()
    assert "Text slide" in pages[1].extract_text()
    # Reuse the retained native PDF so acceptance opens Office exactly once.
    rendered = PresentationRenderer(converter)._rasterize(
        destination, tmp_path / "slides", max_pages=2, max_pixels=4_000_000
    )
    assert not rendered.warnings and len(rendered.assets) == 2
    assert len({asset.sha256 for asset in rendered.assets}) == 2
    quiz, images, provenance = export_fixture(tmp_path)
    quiz = replace(
        quiz,
        questions=tuple(
            replace(
                question,
                stem=f"Slide {index}: which colored shape is shown?",
                choices=tuple(
                    replace(choice, text=label)
                    for choice, label in zip(question.choices, (correct, incorrect), strict=True)
                ),
                rationale=f"Slide {index} shows a {correct.lower()}.",
                learning_objective="lo-1: Identify the colored shape.",
                image_ref=replace(question.image_ref, description=correct),
            )
            for index, (question, correct, incorrect) in enumerate(
                zip(
                    quiz.questions,
                    ("Red rectangle", "Blue oval"),
                    ("Blue oval", "Red rectangle"),
                    strict=True,
                ),
                1,
            )
        ),
    )
    for slide_number, (key, asset) in enumerate(zip(images, rendered.assets, strict=True), 1):
        assert asset.path is not None
        assert asset.locator.slide_number == slide_number
        assert asset.origin == "full-slide-render"
        payload = asset.path.read_bytes()
        assert hashlib.sha256(payload).hexdigest() == asset.sha256
        assert sanitize_quiz_image(payload).payload == payload
        with Image.open(asset.path) as image:
            assert image.size == (asset.width, asset.height)
            rgb = image.convert("RGB").tobytes()
            pixels = list(zip(rgb[0::3], rgb[1::3], rgb[2::3], strict=True))
            red = sum(r > 150 and g < 70 and b < 70 for r, g, b in pixels)
            blue = sum(b > 150 and r < 70 and g < 70 for r, g, b in pixels)
            assert (red > 1000 and blue == 0) if slide_number == 1 else (blue > 1000 and red == 0)
        images[key] = asset.path
        provenance["image_sha256"][key] = asset.sha256
    raw, bundle, pdf = export_reviewed_quiz(quiz, images, provenance, tmp_path / "exports")
    assert parse_native_quiz(raw.read_text(encoding="utf-8")) == quiz
    with ZipFile(bundle) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert archive.read("quiz.json") == raw.read_bytes()
        assert manifest["accepted_payload_sha256"] == hashlib.sha256(raw.read_bytes()).hexdigest()
        assert manifest["provenance"] == provenance
        for key, image_path in images.items():
            assert archive.read(manifest["images"][key]["path"]) == image_path.read_bytes()
    text = "\n".join(page.extract_text() for page in PdfReader(pdf).pages)
    for question in quiz.questions:
        assert question.stem in text and question.rationale in text
        assert grade_answer(quiz, question.id, question.choices[0].id).correct
        assert not grade_answer(quiz, question.id, question.choices[1].id).correct
    assert "Slide 1" in text and "Slide 2" in text
